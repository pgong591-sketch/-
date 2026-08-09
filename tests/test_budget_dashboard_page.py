from pathlib import Path
import inspect

import pandas as pd
from openpyxl import Workbook

import app


def test_budget_actual_period_months_use_pl_detail_ytd_and_latest_month(monkeypatch):
    queries: list[str] = []

    def fake_execute_sql(sql, params=None):
        queries.append(sql)
        return pd.DataFrame({"period": ["202604", "202605", "202606"]})

    monkeypatch.setattr(app, "execute_sql", fake_execute_sql)
    app._budget_actual_period_months_cached.clear()

    months = app._budget_actual_period_months_cached("db", (("db", 1, 1, 1),))

    assert months == {"2026": ["04", "05", "06"]}
    assert queries
    assert "FROM pl_detail" in queries[0]
    assert "ytd_amount IS NOT NULL" in queries[0]
    assert "income_statement" not in queries[0]


def test_budget_actual_period_months_refresh_when_db_signature_changes(monkeypatch):
    calls = []

    def fake_execute_sql(sql, params=None):
        calls.append(sql)
        if len(calls) == 1:
            return pd.DataFrame({"period": ["202606"]})
        return pd.DataFrame({"period": ["202606", "202607"]})

    monkeypatch.setattr(app, "execute_sql", fake_execute_sql)
    app._budget_actual_period_months_cached.clear()

    first = app._budget_actual_period_months_cached("db", (("db", 1, 1, 1),))
    second = app._budget_actual_period_months_cached("db", (("db", 1, 2, 1),))

    assert first == {"2026": ["06"]}
    assert second == {"2026": ["06", "07"]}
    assert len(calls) == 2


def test_budget_month_state_defaults_to_latest_and_preserves_manual_rerun():
    months_by_year = {"2026": ["04", "05", "06"]}

    year, month, options, has_actual = app._resolve_budget_year_month_state(
        ["2026"], months_by_year, None, None, "首页", None
    )
    assert (year, month, options, has_actual) == ("2026", "06", ["04", "05", "06"], True)

    year, month, options, has_actual = app._resolve_budget_year_month_state(
        ["2026"], months_by_year, "2026", "04", "全面预算", "2026"
    )
    assert (year, month, options, has_actual) == ("2026", "04", ["04", "05", "06"], True)


def test_budget_month_default_ignores_later_non_pl_detail_periods():
    months_by_year = {"2026": ["06"]}

    year, month, options, has_actual = app._resolve_budget_year_month_state(
        ["2026"], months_by_year, "2026", None, "首页", None
    )

    assert (year, month, options, has_actual) == ("2026", "06", ["06"], True)


def test_budget_month_state_resets_invalid_month_and_year_change():
    months_by_year = {"2026": ["04", "05", "06"], "2025": ["12"]}

    assert app._resolve_budget_year_month_state(
        ["2026", "2025"], months_by_year, "2026", "03", "全面预算", "2026"
    )[1] == "06"
    assert app._resolve_budget_year_month_state(
        ["2026", "2025"], months_by_year, "2025", "06", "全面预算", "2026"
    )[1] == "12"
    assert app._resolve_budget_year_month_state(
        ["2027"], months_by_year, "2027", "06", "全面预算", "2026"
    ) == ("2027", "03", [f"{idx:02d}" for idx in range(1, 13)], False)


def test_budget_dashboard_uses_budget_actual_month_helper_not_other_reports():
    source = inspect.getsource(app.render_budget_dashboard)

    assert "_budget_actual_period_months()" in source
    assert '_get_year_month_options("income_statement")' not in source


def _write_budget_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2025年总预算（各部门提交）"
    ws.append(["2026年多维集团损益预算"])
    ws.append(["项目", "全年收入", "人工", "净利润", "其他费用"])
    ws.append(["东莞素质中心", 1200.0, "#REF!", 240.0, "#REF!"])
    ws.append(["莞城小学部", 100.0, "#REF!", 20.0, "#REF!"])
    ws.append(["管理中心", 300.0, "#REF!", -30.0, "#REF!"])
    ws.append(["合计", 1500.0, "#REF!", 210.0, "#REF!"])
    ws.append([None, None, None, None, None])
    wb.save(path)
    return path


def _write_budget_workbook_with_quality_center_sheet(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2025年总预算（各部门提交）"
    ws.append(["项目", "全年收入", "净利润"])
    ws.append(["东莞素质中心", 1200.0, 240.0])
    quality = wb.create_sheet("2026年素质中心目标校区四季度营收(含个性化)")
    quality.append([
        "校区", "寒假", "春季", "暑假", "秋季", "个性化", "物品", None,
        "校区", "寒假", "春季", "暑假", "秋季", "合并收入", None, "差额",
    ])
    quality.append(["左侧校区", 1, 2, 3, 4, 5, 6, None, "莞城小学部", 1, 2, 3, 4, 123456.0, None, None])
    quality.append(["左侧校区2", 1, 2, 3, 4, 5, 6, None, "莞城初中部", 1, 2, 3, 4, 234567.0, None, None])
    wb.save(path)
    return path


def test_budget_plan_reads_only_core_fields(tmp_path):
    plan = app.read_budget_plan(_write_budget_workbook(tmp_path / "budget.xlsx"))

    assert plan.columns.tolist() == ["module", "unit_name", "income_budget", "profit_budget", "budget_level"]
    assert "人工" not in plan.columns
    assert "其他费用" not in plan.columns
    assert plan.loc[plan["unit_name"] == "东莞素质中心", "income_budget"].iloc[0] == 1200.0
    assert plan.loc[plan["unit_name"] == "东莞素质中心", "profit_budget"].iloc[0] == 240.0
    assert plan.loc[plan["unit_name"] == "莞城小学部", "module"].iloc[0] == "东莞素质中心"


def test_budget_time_progress_for_march_is_25_percent():
    assert app.budget_time_progress("03") == 0.25
    assert app.budget_time_progress("3月") == 0.25


def test_budget_overview_compares_completion_with_time_progress(tmp_path):
    plan = app.read_budget_plan(_write_budget_workbook(tmp_path / "budget.xlsx"))
    actual = pd.DataFrame(
        [
            {
                "module": "东莞素质中心",
                "unit_name": "莞城小学部",
                "company_code": "101010120",
                "income_actual": 300.0,
                "profit_actual": 60.0,
            },
            {
                "module": "管理中心",
                "unit_name": "管理中心",
                "company_code": "101",
                "income_actual": 50.0,
                "profit_actual": -20.0,
            },
        ]
    )

    overview, detail = app.build_budget_completion_data(plan, actual, "03")

    quality = overview.loc[overview["模块名称"] == "东莞素质中心"].iloc[0]
    assert quality["收入预算"] == 1200.0
    assert quality["收入完成率"] == 0.25
    assert quality["时间进度"] == 0.25
    assert quality["状态"] == "正常"

    management = overview.loc[overview["模块名称"] == "管理中心"].iloc[0]
    assert management["收入完成率"] < management["时间进度"]
    assert management["状态"] == "滞后"

    drill = detail.loc[detail["公司名称"] == "莞城小学部"].iloc[0]
    assert drill["收入预算"] == 100.0
    assert drill["利润预算"] == 20.0
    assert drill["收入实际"] == 300.0


def test_budget_overview_applies_internal_adjustment_to_module_and_total_only():
    plan = pd.DataFrame(
        [
            {"module": "东莞素质中心", "unit_name": "东莞素质中心", "income_budget": 1000.0, "profit_budget": 100.0, "budget_level": "module"},
            {"module": "管理中心", "unit_name": "管理中心", "income_budget": 500.0, "profit_budget": 50.0, "budget_level": "module"},
            {"module": "合计", "unit_name": "合计", "income_budget": 1500.0, "profit_budget": 150.0, "budget_level": "module"},
        ]
    )
    actual = pd.DataFrame(
        [
            {"module": "东莞素质中心", "unit_name": "莞城小学部", "company_code": "101010120", "income_actual": 300.0, "profit_actual": 60.0},
            {"module": "管理中心", "unit_name": "管理中心", "company_code": "101", "income_actual": 50.0, "profit_actual": -20.0},
        ]
    )
    actual.attrs["budget_internal_adjustments"] = {
        "module_adjustments": {"东莞素质中心": 40.0},
        "total_adjustment": 70.0,
    }

    overview, detail = app.build_budget_completion_data(plan, actual, "03")

    quality = overview.loc[overview["模块名称"] == "东莞素质中心"].iloc[0]
    total = overview.loc[overview["模块名称"] == "合计"].iloc[0]
    drill = detail.loc[detail["公司名称"] == "莞城小学部"].iloc[0]
    assert quality["收入实际"] == 260.0
    assert quality["利润实际"] == 60.0
    assert total["收入实际"] == 280.0
    assert total["利润实际"] == 40.0
    assert drill["收入实际"] == 300.0


def test_budget_completion_filters_pseudo_detail_rows_and_keeps_single_total():
    plan = pd.DataFrame(
        [
            {"module": "青少年宫", "unit_name": "青少年宫", "income_budget": 100.0, "profit_budget": 10.0, "budget_level": "module"},
            {"module": "青少年宫", "unit_name": "真实未分组公司", "income_budget": 20.0, "profit_budget": 2.0, "budget_level": "unit"},
            {"module": "青少年宫", "unit_name": "合计", "income_budget": 999.0, "profit_budget": 999.0, "budget_level": "unit"},
            {"module": "合并", "unit_name": "合并", "income_budget": 888.0, "profit_budget": 888.0, "budget_level": "module"},
            {"module": "合计", "unit_name": "合计", "income_budget": 777.0, "profit_budget": 777.0, "budget_level": "module"},
        ]
    )
    actual = pd.DataFrame(
        [
            {"module": "青少年宫", "unit_name": "真实未分组公司", "company_code": "1011801", "income_actual": 30.0, "profit_actual": 3.0},
            {"module": "青少年宫", "unit_name": "合并", "company_code": "SUMMARY_10118", "income_actual": 999.0, "profit_actual": 999.0},
            {"module": "SUMMARY_收入合计", "unit_name": "SUMMARY_收入合计", "company_code": "SUMMARY_ROW", "income_actual": 888.0, "profit_actual": 888.0},
        ]
    )

    overview, detail = app.build_budget_completion_data(plan, actual, "03")

    assert overview["模块名称"].tolist().count("合计") == 1
    assert "合并" not in set(overview["模块名称"])
    assert "合计" not in set(detail["公司名称"])
    assert "合并" not in set(detail["公司名称"])
    assert "真实未分组公司" in set(detail["公司名称"])
    total = overview.loc[overview["模块名称"] == "合计"].iloc[0]
    assert total["收入预算"] == 100.0
    assert total["收入实际"] == 30.0


def test_budget_bridge_note_explains_raw_drilldown_and_adjusted_main_value():
    detail = pd.DataFrame(
        [
            {"module": "东莞素质中心", "公司名称": "莞城小学部", "收入实际": 300.0, "利润实际": 60.0},
            {"module": "东莞素质中心", "公司名称": "万江校区", "收入实际": 200.0, "利润实际": 40.0},
        ]
    )
    detail.attrs["budget_internal_adjustments"] = {
        "module_adjustments": {"东莞素质中心": 50.0},
        "total_adjustment": 50.0,
    }

    bridge = app._budget_module_bridge(detail, "东莞素质中心")
    html = app._budget_bridge_note_html("东莞素质中心", bridge)

    assert bridge["raw_income"] == 500.0
    assert bridge["income_adjustment"] == 50.0
    assert bridge["adjusted_income"] == 450.0
    assert "下钻为单体原始本年累计" in html
    assert "减内部抵消" in html


def test_load_budget_actuals_uses_pl_detail_ytd_and_preferred_rows(monkeypatch):
    queries: list[str] = []
    rows = pd.DataFrame(
        [
            {"id": 1, "公司编码": "1011801", "company_code": "1011801", "item_code": "SUMMARY_收入合计", "item_name": app.PL_REVENUE_ITEM, "amount": 100.0, "ytd_amount": 100.0, "short_name": "莞城青少年宫", "company_name": "莞城青少年宫", "parent_code": "10118", "business_group": "", "公司": "莞城青少年宫"},
            {"id": 2, "公司编码": "1011801", "company_code": "1011801", "item_code": "OPERATING_收入合计", "item_name": app.PL_REVENUE_ITEM, "amount": 120.0, "ytd_amount": 120.0, "short_name": "莞城青少年宫", "company_name": "莞城青少年宫", "parent_code": "10118", "business_group": "", "公司": "莞城青少年宫"},
            {"id": 3, "公司编码": "1011801", "company_code": "1011801", "item_code": "5401", "item_name": app.PL_REVENUE_ITEM, "amount": 999.0, "ytd_amount": 999.0, "short_name": "莞城青少年宫", "company_name": "莞城青少年宫", "parent_code": "10118", "business_group": "", "公司": "莞城青少年宫"},
            {"id": 4, "公司编码": "1011801", "company_code": "1011801", "item_code": "OPERATING_净利润", "item_name": app.PL_NET_PROFIT_ITEM, "amount": 30.0, "ytd_amount": 30.0, "short_name": "莞城青少年宫", "company_name": "莞城青少年宫", "parent_code": "10118", "business_group": "", "公司": "莞城青少年宫"},
            {"id": 5, "公司编码": "1011801", "company_code": "1011801", "item_code": "OPERATING_成本费用合计", "item_name": app.PL_COST_TOTAL_ITEM, "amount": 90.0, "ytd_amount": 90.0, "short_name": "莞城青少年宫", "company_name": "莞城青少年宫", "parent_code": "10118", "business_group": "", "公司": "莞城青少年宫"},
        ]
    )

    def fake_execute_sql(sql, params=None):
        queries.append(sql)
        return rows.copy()

    monkeypatch.setattr(app, "execute_sql", fake_execute_sql)
    monkeypatch.setattr(app, "get_consolidation_company_codes", lambda code: {"10118": ["1011801"]}.get(str(code), [str(code)]))

    actual = app.load_budget_actuals("202603")

    assert all("income_statement" not in query for query in queries)
    assert "COALESCE(d.ytd_amount, 0) AS amount" in inspect.getsource(app._budget_pl_detail_ytd_source_rows)
    row = actual.iloc[0]
    assert row["income_actual"] == 120.0
    assert row["profit_actual"] == 30.0


def test_budget_ytd_internal_fee_split_does_not_assign_unmatched_residual_to_101(monkeypatch):
    monkeypatch.setattr(
        app,
        "get_consolidation_company_codes",
        lambda code: {
            "10101": ["10101", "1010101", "101010102"],
            "10204": ["10204", "1020401"],
        }.get(str(code), [str(code)]),
    )
    monkeypatch.setattr(app, "_operating_card_company_name_map", lambda codes: {str(code): str(code) for code in codes})
    rows = pd.DataFrame(
        [
            {"id": 1, "company_code": "1010101", "公司": "东莞非学科管理中心", "item_code": "OPERATING_成本费用合计", "item_name": app.PL_COST_TOTAL_ITEM, "amount": 0.0},
            {"id": 2, "company_code": "101010102", "公司": "华凯校区", "item_code": "560216", "item_name": app.HOME_MANAGEMENT_FEE_ITEM, "amount": 100.0},
            {"id": 3, "company_code": "101010102", "公司": "华凯校区", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 1000.0},
            {"id": 4, "company_code": "1020401", "公司": "莞城书馆", "item_code": "560216", "item_name": app.HOME_MANAGEMENT_FEE_ITEM, "amount": 30.0},
            {"id": 5, "company_code": "1020401", "公司": "莞城书馆", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 300.0},
            {"id": 6, "company_code": "10204", "公司": "书馆管理中心", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 0.0},
            {"id": 7, "company_code": "1011801", "公司": "青少年宫", "item_code": "560216", "item_name": app.HOME_MANAGEMENT_FEE_ITEM, "amount": 5.0},
            {"id": 8, "company_code": "1011801", "公司": "青少年宫", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 50.0},
        ]
    )
    metrics = app._operating_card_company_metrics_from_source(rows)

    split = app._operating_card_internal_fee_split(
        rows,
        metrics,
        ["1010101", "101010102", "10204", "1020401", "1011801"],
        assign_residual_to_management=False,
    )

    assert split["non_subject_fee"].get("101010102", 0.0) == 0.0
    assert split["non_subject_unmatched_fee"]["101010102"] == 100.0
    assert split["unmatched_eryu_fee"]["1020401"] == 30.0
    assert split["management_fee"].get("101010102", 0.0) == 0.0
    assert split["management_fee"].get("1020401", 0.0) == 0.0
    assert split["management_fee"]["1011801"] == 5.0


def test_budget_ytd_scope_only_offsets_when_payer_and_receiver_are_in_scope(monkeypatch):
    monkeypatch.setattr(
        app,
        "get_consolidation_company_codes",
        lambda code: {
            "10101": ["10101", "1010101", "101010102"],
            "10204": ["10204", "1020401"],
            "10118": ["10118", "1011801"],
        }.get(str(code), [str(code)]),
    )
    monkeypatch.setattr(app, "_operating_card_company_name_map", lambda codes: {str(code): str(code) for code in codes})
    rows = pd.DataFrame(
        [
            {"company_code": "101010102", "公司": "华凯校区", "item_code": "OPERATING_收入合计", "item_name": app.PL_REVENUE_ITEM, "amount": 1_080_082.99},
            {"company_code": "101010102", "公司": "华凯校区", "item_code": "OPERATING_净利润", "item_name": app.PL_NET_PROFIT_ITEM, "amount": 100_000.0},
            {"company_code": "101010102", "公司": "华凯校区", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 1_080_082.99},
            {"company_code": "101010102", "公司": "华凯校区", "item_code": "560216", "item_name": app.HOME_MANAGEMENT_FEE_ITEM, "amount": 88_679.98},
            {"company_code": "1010101", "公司": "东莞非学科管理中心", "item_code": "OPERATING_收入合计", "item_name": app.PL_REVENUE_ITEM, "amount": 0.0},
            {"company_code": "1010101", "公司": "东莞非学科管理中心", "item_code": "OPERATING_净利润", "item_name": app.PL_NET_PROFIT_ITEM, "amount": 0.0},
            {"company_code": "1010101", "公司": "东莞非学科管理中心", "item_code": "OPERATING_成本费用合计", "item_name": app.PL_COST_TOTAL_ITEM, "amount": 0.0},
            {"company_code": "1020401", "公司": "莞城书馆", "item_code": "OPERATING_收入合计", "item_name": app.PL_REVENUE_ITEM, "amount": 2_148_024.47},
            {"company_code": "1020401", "公司": "莞城书馆", "item_code": "OPERATING_净利润", "item_name": app.PL_NET_PROFIT_ITEM, "amount": 200_000.0},
            {"company_code": "1020401", "公司": "莞城书馆", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 2_148_024.47},
            {"company_code": "1020401", "公司": "莞城书馆", "item_code": "560216", "item_name": app.HOME_MANAGEMENT_FEE_ITEM, "amount": 198_641.78},
            {"company_code": "10204", "公司": "书馆管理中心", "item_code": "OPERATING_收入合计", "item_name": app.PL_REVENUE_ITEM, "amount": 0.0},
            {"company_code": "10204", "公司": "书馆管理中心", "item_code": "OPERATING_净利润", "item_name": app.PL_NET_PROFIT_ITEM, "amount": 0.0},
            {"company_code": "10204", "公司": "书馆管理中心", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 0.0},
            {"company_code": "1011801", "公司": "青少年宫", "item_code": "OPERATING_收入合计", "item_name": app.PL_REVENUE_ITEM, "amount": 703_698.77},
            {"company_code": "1011801", "公司": "青少年宫", "item_code": "OPERATING_净利润", "item_name": app.PL_NET_PROFIT_ITEM, "amount": 70_000.0},
            {"company_code": "1011801", "公司": "青少年宫", "item_code": "540101", "item_name": app.HOME_MAIN_REVENUE_ITEM, "amount": 703_698.77},
        ]
    )

    def fake_source_rows(period, company_codes=None):
        result = rows.copy()
        if company_codes is not None:
            wanted = {str(code) for code in company_codes}
            result = result[result["company_code"].astype(str).isin(wanted)].copy()
        for column, default in {
            "id": 1,
            "公司编码": "",
            "ytd_amount": 0.0,
            "short_name": "",
            "company_name": "",
            "parent_code": "",
            "business_group": "",
        }.items():
            if column not in result.columns:
                result[column] = default
        result["公司编码"] = result["company_code"]
        result["ytd_amount"] = result["amount"]
        result["short_name"] = result["公司"]
        result["company_name"] = result["公司"]
        return result

    def scoped_income(codes: list[str]) -> tuple[float, dict]:
        actual_rows = []
        source = fake_source_rows("202603", codes)
        metrics = app._operating_card_company_metrics_from_source(source)
        for row in metrics.to_dict("records"):
            actual_rows.append(
                {
                    "module": app._budget_module_for_actual_company(row["company_code"], row["公司"]),
                    "unit_name": row["公司"],
                    "company_code": row["company_code"],
                    "income_actual": row["收入"],
                    "profit_actual": row["净利润"],
                }
            )
        actual = pd.DataFrame(actual_rows, columns=app._budget_empty_actual_frame().columns)
        monkeypatch.setattr(app, "_budget_pl_detail_ytd_source_rows", fake_source_rows)
        actual = app._budget_attach_internal_adjustments(actual, "202603")
        plan = pd.DataFrame(
            [
                {"module": "东莞素质中心", "unit_name": "东莞素质中心", "income_budget": 9_999_999.0, "profit_budget": 1.0, "budget_level": "module"},
                {"module": "尔遇书馆", "unit_name": "尔遇书馆", "income_budget": 9_999_999.0, "profit_budget": 1.0, "budget_level": "module"},
                {"module": "青少年宫", "unit_name": "青少年宫", "income_budget": 9_999_999.0, "profit_budget": 1.0, "budget_level": "module"},
                {"module": "合计", "unit_name": "合计", "income_budget": 9_999_999.0, "profit_budget": 1.0, "budget_level": "module"},
            ]
        )
        overview, _ = app.build_budget_completion_data(plan, actual, "03")
        return (
            app._safe_float(overview.loc[overview["模块名称"] == "合计", "收入实际"].iloc[0]),
            actual.attrs["budget_internal_adjustments"],
        )

    assert round(scoped_income(["101010102"])[0], 2) == 1_080_082.99
    assert round(scoped_income(["1020401"])[0], 2) == 2_148_024.47
    assert round(scoped_income(["1011801"])[0], 2) == 703_698.77
    assert round(scoped_income(["1010101"])[0], 2) == 0.0
    assert round(scoped_income(["10204"])[0], 2) == 0.0

    non_subject_income, non_subject_context = scoped_income(["101010102", "1010101"])
    eryu_income, eryu_context = scoped_income(["1020401", "10204"])
    assert round(non_subject_income, 2) == 991_403.01
    assert round(non_subject_context["total_adjustment"], 2) == 88_679.98
    assert round(eryu_income, 2) == 1_949_382.69
    assert round(eryu_context["total_adjustment"], 2) == 198_641.78


def test_module_budget_rows_backfill_single_actual_drilldown():
    plan = pd.DataFrame(
        [
            {"module": "青少年宫", "unit_name": "青少年宫", "income_budget": 400.0, "profit_budget": 40.0, "budget_level": "module"},
            {"module": "多维学校", "unit_name": "多维学校", "income_budget": 500.0, "profit_budget": 50.0, "budget_level": "module"},
            {"module": "新阳光幼儿园", "unit_name": "新阳光幼儿园", "income_budget": 600.0, "profit_budget": 60.0, "budget_level": "module"},
            {"module": "托育项目", "unit_name": "托育项目", "income_budget": 700.0, "profit_budget": -70.0, "budget_level": "module"},
            {"module": "尔遇书城", "unit_name": "尔遇书城", "income_budget": 800.0, "profit_budget": -80.0, "budget_level": "module"},
        ]
    )
    actual = pd.DataFrame(
        [
            {"module": "青少年宫", "unit_name": "少年宫", "company_code": "1011801", "income_actual": 100.0, "profit_actual": 10.0},
            {"module": "多维学校", "unit_name": "多维学校", "company_code": "1010201", "income_actual": 110.0, "profit_actual": 11.0},
            {"module": "新阳光幼儿园", "unit_name": "茶山幼儿园", "company_code": "1010701", "income_actual": 120.0, "profit_actual": 12.0},
            {"module": "托育项目", "unit_name": "茶山托育", "company_code": "1010801", "income_actual": 130.0, "profit_actual": -13.0},
            {"module": "尔遇书城", "unit_name": "尔遇书城", "company_code": "1010601", "income_actual": 140.0, "profit_actual": -14.0},
        ]
    )

    _, detail = app.build_budget_completion_data(plan, actual, "03")

    expected = {
        "青少年宫": ("少年宫", 400.0),
        "多维学校": ("多维学校", 500.0),
        "新阳光幼儿园": ("茶山幼儿园", 600.0),
        "托育项目": ("茶山托育", 700.0),
        "尔遇书城": ("尔遇书城", 800.0),
    }
    for module, (name, income_budget) in expected.items():
        view = app._budget_module_drilldown_view(module, detail, app.budget_time_progress("03"))
        assert len(view) == 1
        assert view["公司名称"].iloc[0] == name
        assert view["收入预算"].iloc[0] == income_budget
        assert view["收入实际"].iloc[0] > 0


def test_budget_actuals_drop_stale_parent_duplicate_before_drilldown():
    actual = pd.DataFrame(
        [
            {
                "module": "新阳光幼儿园",
                "unit_name": "幼儿园",
                "company_code": "10107",
                "parent_code": "101",
                "income_actual": 120.0,
                "profit_actual": -12.0,
            },
            {
                "module": "新阳光幼儿园",
                "unit_name": "东莞市茶山新阳光幼儿园",
                "company_code": "1010702",
                "parent_code": "10107",
                "income_actual": 120.0,
                "profit_actual": -12.0,
            },
            {
                "module": "托育项目",
                "unit_name": "茶山托育项目",
                "company_code": "1010703",
                "parent_code": "10107",
                "income_actual": 80.0,
                "profit_actual": -8.0,
            },
        ]
    )

    pruned = app._budget_prune_stale_parent_actuals(actual)

    assert pruned["company_code"].tolist() == ["1010702", "1010703"]


def test_module_budget_only_drilldown_uses_pending_actual_not_zero():
    plan = pd.DataFrame(
        [
            {"module": "青少年宫", "unit_name": "青少年宫", "income_budget": 400.0, "profit_budget": 40.0, "budget_level": "module"},
        ]
    )
    actual = pd.DataFrame(columns=["module", "unit_name", "company_code", "income_actual", "profit_actual"])

    _, detail = app.build_budget_completion_data(plan, actual, "03")
    view = app._budget_module_drilldown_view("青少年宫", detail, app.budget_time_progress("03"))

    assert len(view) == 1
    assert view["公司名称"].iloc[0] == "青少年宫"
    assert view["收入预算"].iloc[0] == 400.0
    assert view["收入实际"].iloc[0] == "待接入"
    assert view["收入进度"].iloc[0] == "待接入"


def test_unclassified_drilldown_expands_actual_names_without_fake_budget():
    actual = pd.DataFrame(
        [
            {"module": "未分组", "unit_name": "东莞国际", "company_code": "A", "income_actual": 100.0, "profit_actual": -10.0},
            {"module": "未分组", "unit_name": "探幽文旅", "company_code": "B", "income_actual": 200.0, "profit_actual": -20.0},
        ]
    )

    _, detail = app.build_budget_completion_data(app._budget_empty_plan_frame(), actual, "03")
    view = app._budget_module_drilldown_view("未分组", detail, app.budget_time_progress("03"))

    assert view["公司名称"].tolist() == ["东莞国际", "探幽文旅"]
    assert view["收入预算"].tolist() == ["暂无预算", "暂无预算"]
    assert view["利润预算"].tolist() == ["暂无预算", "暂无预算"]
    assert view["收入实际"].tolist() == [100.0, 200.0]


def test_negative_profit_budget_loss_overrun_is_not_ahead():
    plan = pd.DataFrame(
        [
            {
                "module": "管理中心",
                "unit_name": "管理中心",
                "income_budget": 1000.0,
                "profit_budget": -100.0,
                "budget_level": "module",
            }
        ]
    )
    actual = pd.DataFrame(
        [
            {
                "module": "管理中心",
                "unit_name": "管理中心",
                "company_code": "101",
                "income_actual": 250.0,
                "profit_actual": -200.0,
            }
        ]
    )

    overview, _ = app.build_budget_completion_data(plan, actual, "03")
    row = overview.iloc[0]

    assert pd.isna(row["利润完成率"])
    assert row["利润进度差"] == -1.75
    assert row["状态"] == "滞后"


def test_negative_profit_budget_lower_loss_can_be_ahead():
    plan = pd.DataFrame(
        [
            {
                "module": "管理中心",
                "unit_name": "管理中心",
                "income_budget": 1000.0,
                "profit_budget": -100.0,
                "budget_level": "module",
            }
        ]
    )
    actual = pd.DataFrame(
        [
            {
                "module": "管理中心",
                "unit_name": "管理中心",
                "company_code": "101",
                "income_actual": 250.0,
                "profit_actual": -10.0,
            }
        ]
    )

    overview, _ = app.build_budget_completion_data(plan, actual, "03")
    row = overview.iloc[0]

    assert pd.isna(row["利润完成率"])
    assert row["利润进度差"] == 0.15
    assert row["状态"] == "超前"


def test_budget_default_path_and_missing_message_do_not_leak_wechat_temp_path():
    default_path = str(app.BUDGET_WORKBOOK_PATH)
    message = app._budget_missing_file_message()

    assert "xwechat_files" not in default_path
    assert "RWTemp" not in default_path
    assert default_path.endswith("data/2026年集团预算.xlsx")
    assert "xwechat_files" not in message
    assert "RWTemp" not in message


def test_budget_navigation_entry_is_under_business_center():
    entries = app.NAV_MODULE_SECTIONS["经营中心"]["经营看板"]

    assert "全面预算" in entries
    assert "利润表总览驾驶舱" not in entries
    assert entries.index("全面预算") == 1
    assert app.NAV_LABELS["全面预算"] == "全面预算"


def test_budget_overview_view_keeps_time_progress_inside_income_and_profit_columns(tmp_path):
    plan = app.read_budget_plan(_write_budget_workbook(tmp_path / "budget.xlsx"))
    actual = pd.DataFrame(
        [
            {
                "module": "东莞素质中心",
                "unit_name": "莞城小学部",
                "company_code": "101010120",
                "income_actual": 300.0,
                "profit_actual": 60.0,
            }
        ]
    )
    overview, detail = app.build_budget_completion_data(plan, actual, "03")

    overview_view = app._budget_overview_view(overview)
    detail_view = app._budget_detail_view(detail, app.budget_time_progress("03"))

    assert "时间进度" not in overview_view.columns
    assert "收入进度" in overview_view.columns
    assert "利润进度" in overview_view.columns
    assert "时间进度 25.00%" in overview_view.loc[overview_view["模块名称"] == "东莞素质中心", "收入进度"].iloc[0]
    assert "时间进度 25.00%" in overview_view.loc[overview_view["模块名称"] == "东莞素质中心", "利润进度"].iloc[0]
    assert "时间进度" not in detail_view.columns


def test_budget_kpi_cards_render_final_two_card_dashboard(tmp_path):
    plan = app.read_budget_plan(_write_budget_workbook(tmp_path / "budget.xlsx"))
    actual = pd.DataFrame(
        [
            {
                "module": "东莞素质中心",
                "unit_name": "莞城小学部",
                "company_code": "101010120",
                "income_actual": 300.0,
                "profit_actual": 60.0,
            }
        ]
    )
    overview, _ = app.build_budget_completion_data(plan, actual, "03")

    html = app._budget_kpi_cards_html(overview, "03", app.budget_time_progress("03"))

    assert html.count('class="budget-kpi-card"') == 2
    assert "收入预算完成情况" in html
    assert "利润预算完成情况" in html
    assert "时间进度 / 平均完成度" not in html
    assert "实际收入" in html
    assert "实际利润" in html
    assert "时间进度 25.0%" in html
    assert "平均完成度" in html
    assert "状态：" in html
    assert "&lt;div" not in html


def test_budget_kpi_cards_show_average_completion_excluding_total_and_missing_values():
    overview = pd.DataFrame(
        [
            {"模块名称": "东莞素质中心", "收入预算": 100.0, "收入实际": 20.0, "收入完成率": 0.2, "利润预算": 10.0, "利润实际": 6.0, "利润完成率": 0.6, "状态": "正常"},
            {"模块名称": "管理中心", "收入预算": 100.0, "收入实际": 40.0, "收入完成率": 0.4, "利润预算": -10.0, "利润实际": -4.0, "利润完成率": None, "状态": "滞后"},
            {"模块名称": "无预算", "收入预算": 0.0, "收入实际": 10.0, "收入完成率": None, "利润预算": 0.0, "利润实际": 1.0, "利润完成率": None, "状态": "暂无预算"},
            {"模块名称": "合计", "收入预算": 999.0, "收入实际": 999.0, "收入完成率": 9.9, "利润预算": 999.0, "利润实际": 999.0, "利润完成率": 9.9, "状态": "正常"},
        ]
    )

    assert round(app._budget_average_completion(overview, "income"), 4) == 0.3
    assert round(app._budget_average_completion(overview, "profit"), 4) == 0.6

    html = app._budget_kpi_cards_html(overview, "03", app.budget_time_progress("03"))

    assert "时间进度 25.0% · 平均完成度" in html
    assert "30.0%" in html
    assert "60.0%" in html


def test_budget_profit_average_completion_shows_no_comparable_units():
    overview = pd.DataFrame(
        [
            {"模块名称": "亏损预算", "收入预算": 100.0, "收入实际": 30.0, "收入完成率": 0.3, "利润预算": -10.0, "利润实际": -2.0, "利润完成率": None, "状态": "正常"},
            {"模块名称": "零利润预算", "收入预算": 50.0, "收入实际": 10.0, "收入完成率": 0.2, "利润预算": 0.0, "利润实际": 1.0, "利润完成率": None, "状态": "暂无预算"},
            {"模块名称": "合计", "收入预算": 150.0, "收入实际": 40.0, "收入完成率": 0.27, "利润预算": -10.0, "利润实际": -1.0, "利润完成率": None, "状态": "正常"},
        ]
    )

    assert app._budget_average_completion(overview, "profit") is None

    html = app._budget_kpi_cards_html(overview, "03", app.budget_time_progress("03"))

    assert "暂无可比单位" in html


def test_budget_comparison_view_matches_target_columns(tmp_path):
    plan = app.read_budget_plan(_write_budget_workbook(tmp_path / "budget.xlsx"))
    actual = pd.DataFrame(
        [
            {
                "module": "东莞素质中心",
                "unit_name": "莞城小学部",
                "company_code": "101010120",
                "income_actual": 300.0,
                "profit_actual": 60.0,
            }
        ]
    )
    overview, _ = app.build_budget_completion_data(plan, actual, "03")

    view = app._budget_comparison_table_view(overview)

    assert view.columns.tolist() == [
        "经营单位",
        "收入预算",
        "收入实际",
        "收入完成率",
        "利润预算",
        "利润实际",
        "利润完成率",
        "时间进度",
        "进度判断",
    ]
    assert view.loc[view["经营单位"] == "东莞素质中心", "收入预算"].iloc[0] == 0.12
    assert view.loc[view["经营单位"] == "东莞素质中心", "时间进度"].iloc[0] == 0.25


def test_quality_center_campus_targets_read_column_n_total_target(tmp_path):
    workbook = _write_budget_workbook_with_quality_center_sheet(tmp_path / "budget.xlsx")

    targets = app.read_quality_center_campus_budget_targets(workbook)

    assert targets.columns.tolist() == ["campus_name", "income_budget"]
    assert targets["campus_name"].tolist() == ["莞城小学部", "莞城初中部"]
    assert targets["income_budget"].tolist() == [123456.0, 234567.0]


def test_quality_center_drilldown_view_matches_actual_income_and_keeps_unmatched_pending(tmp_path):
    workbook = _write_budget_workbook_with_quality_center_sheet(tmp_path / "budget.xlsx")
    targets = app.read_quality_center_campus_budget_targets(workbook)
    actual = pd.DataFrame(
        [
            {"actual_name": "莞城小学部", "company_code": "101010120", "actual_income": 61728.0},
            {"actual_name": "莞城初中部", "company_code": "101010128", "actual_income": 100000.0},
        ]
    )

    view = app._budget_quality_center_drilldown_view(targets, app.budget_time_progress("03"), actual)

    assert view.columns.tolist() == ["校区名称", "全年收入预算目标", "实际收入", "收入完成率", "时间进度", "进度差", "状态"]
    assert view.loc[view["校区名称"] == "莞城小学部", "全年收入预算目标"].iloc[0] == 123456.0
    assert view.loc[view["校区名称"] == "莞城小学部", "实际收入"].iloc[0] == 61728.0
    assert view.loc[view["校区名称"] == "莞城小学部", "收入完成率"].iloc[0] == 0.5
    assert view.loc[view["校区名称"] == "莞城小学部", "时间进度"].iloc[0] == 0.25


def test_quality_center_match_statuses_and_unmatched_list():
    targets = pd.DataFrame(
        [
            {"campus_name": "万江校区", "income_budget": 100000.0},
            {"campus_name": "初中总部校区", "income_budget": 200000.0},
            {"campus_name": "未知校区", "income_budget": 300000.0},
        ]
    )
    actual = pd.DataFrame(
        [
            {"actual_name": "万江", "company_code": "101010124", "actual_income": 50000.0},
            {"actual_name": "莞城初中部", "company_code": "101010128", "actual_income": 80000.0},
        ]
    )

    matched = app.match_quality_center_campus_actuals(targets, actual)
    view = app._budget_quality_center_drilldown_view(targets, 0.25, actual)
    issues = app.quality_center_unmatched_items(matched)

    assert matched.loc[matched["campus_name"] == "万江校区", "match_status"].iloc[0] == "已匹配"
    assert view.loc[view["校区名称"] == "万江校区", "实际收入"].iloc[0] == 50000.0
    assert matched.loc[matched["campus_name"] == "初中总部校区", "match_status"].iloc[0] == "已匹配"
    assert view.loc[view["校区名称"] == "初中总部校区", "实际收入"].iloc[0] == 80000.0
    assert view.loc[view["校区名称"] == "未知校区", "实际收入"].iloc[0] == "待匹配"
    assert issues["预算校区名称"].tolist() == ["未知校区"]


def test_quality_center_confirmed_mappings_and_special_statuses():
    targets = pd.DataFrame(
        [
            {"campus_name": "南城虎翼营", "income_budget": 100000.0},
            {"campus_name": "茶山校区", "income_budget": 100000.0},
            {"campus_name": "华凯校区", "income_budget": 100000.0},
            {"campus_name": "松山湖校区", "income_budget": 100000.0},
            {"campus_name": "产品中心直营校", "income_budget": 100000.0},
        ]
    )
    actual = pd.DataFrame(
        [
            {"actual_name": "南城虎翼营", "company_code": "101010133", "actual_income": 11000.0},
            {"actual_name": "茶山学前校区", "company_code": "101010131", "actual_income": 22000.0},
            {"actual_name": "南城华凯校区", "company_code": "101010102", "actual_income": 33000.0},
            {"actual_name": "南城", "company_code": "101010133", "actual_income": 33000.0},
        ]
    )

    matched = app.match_quality_center_campus_actuals(targets, actual)
    view = app._budget_quality_center_drilldown_view(targets, 0.25, actual)
    issues = app.quality_center_unmatched_items(matched)

    assert matched.loc[matched["campus_name"] == "南城虎翼营", "actual_name"].iloc[0] == "南城虎翼营"
    assert matched.loc[matched["campus_name"] == "茶山校区", "actual_name"].iloc[0] == "茶山学前校区"
    assert matched.loc[matched["campus_name"] == "华凯校区", "actual_name"].iloc[0] == "南城华凯校区"
    assert matched.loc[matched["campus_name"] == "华凯校区", "company_code"].iloc[0] == "101010102"
    assert set(matched["match_status"]) == {"已匹配", "待开业", "已取消"}
    assert view.loc[view["校区名称"] == "华凯校区", "实际收入"].iloc[0] == 33000.0
    assert view.loc[view["校区名称"] == "松山湖校区", "实际收入"].iloc[0] == "待开业"
    assert view.loc[view["校区名称"] == "松山湖校区", "状态"].iloc[0] == "待开业"
    assert view.loc[view["校区名称"] == "产品中心直营校", "实际收入"].iloc[0] == "已取消"
    assert view.loc[view["校区名称"] == "产品中心直营校", "状态"].iloc[0] == "已取消"
    assert issues.empty


def test_budget_quality_center_mapping_comes_from_base_settings_service():
    import inspect

    source = inspect.getsource(app._budget_quality_confirmed_candidate)

    assert "get_budget_campus_name_mappings" in source
    assert "BUDGET_QUALITY_CENTER_CONFIRMED_MAPPINGS" not in inspect.getsource(app)


def test_budget_comparison_table_uses_inline_links_for_drilldown():
    import inspect

    source = inspect.getsource(app._render_budget_comparison_table)

    assert "budget-module-link" in source
    assert "budget_drill" in source
    assert 'target="_top"' in source
    assert "st.columns" not in source
    assert ".button(" not in source
    assert "_render_budget_comparison_table" in source
    assert "选择经营单位下钻" not in source
    assert "打开下钻明细" not in source


def test_budget_table_cells_have_alignment_lines_and_negative_styles():
    assert app._budget_cell_html(-12.3, "number").startswith('<div class="budget-table-num budget-negative"')
    assert app._budget_cell_html(-0.123, "rate").startswith('<div class="budget-table-num budget-negative"')
    assert "budget-status-lag" in app._budget_cell_html("滞后", "status")

    import inspect

    source = inspect.getsource(app._render_budget_comparison_table)
    assert "单位：万元" in source
    assert "budget-comparison-table" in source
    assert "border-collapse:collapse" in source
    assert "border-top" in source
    assert "border-right" in source


def test_budget_main_and_drill_tables_have_sticky_header_and_first_column():
    import inspect

    main_source = inspect.getsource(app._render_budget_comparison_table)
    drill_source = inspect.getsource(app._budget_drill_table_html)

    assert ".budget-table-scroll" in main_source
    assert "overflow:auto" in main_source
    assert ".budget-comparison-table th{{position:sticky;top:0;z-index:4" in main_source
    assert ".budget-comparison-table th:first-child{{left:0;z-index:7" in main_source
    assert ".budget-table-first{{position:sticky;left:0;z-index:3" in main_source
    assert ".budget-drill-wrap{{margin-top:10px;overflow:auto" in drill_source
    assert ".budget-drill-table th{{position:sticky;top:0;z-index:4" in drill_source
    assert ".budget-drill-table th:first-child{{left:0;z-index:7" in drill_source
    assert ".budget-drill-table td:first-child{{position:sticky;left:0;z-index:3" in drill_source


def test_budget_policy_note_is_compact_text_below_main_table():
    note = app._budget_policy_note(0.25)

    assert note.startswith("说明：")
    assert "当前时间进度为 25.00%" in note
    assert "实际数来源为收入成本费用表本年累计" in note
    assert "平均完成度仅统计正预算且有实际数的经营单位" in note
    assert "负利润预算按减亏进度单独判断" in note


def test_budget_page_source_has_no_raw_card_html_leakage():
    import inspect

    source = inspect.getsource(app.render_budget_dashboard)

    assert "focus-expense-card" not in source
    assert "_render_budget_kpi_cards" in source
    assert "_budget_policy_note" in source
    assert "_render_budget_policy_card" not in source
    assert "_render_budget_comparison_table" in source
    assert "选择经营单位下钻" not in source
    assert "打开下钻明细" not in source
    assert "st.dataframe" in inspect.getsource(app._render_budget_dataframe)
    assert "st.info(f\"{int(selected_month)}月时间进度" not in source
